"""Lifecycle: what `delete_promoted()` promises, and proof that it keeps that promise.

The inventory of derived tables is built from live ORM metadata rather than from a list in this file,
because the defect this suite exists to prevent is precisely a list drifting out of step with the
models.  Row snapshots are compared, not counts: a delete that removes one row and adds another in
the same table has the same count and is still wrong.
"""

from __future__ import annotations

from sqlalchemy import select
from tests.fixtures.fieldops import fetch, ingest_v71, ingest_v72, promote_file, reingest

from drilling_intelligence.core.enums import KnowledgeOrigin
from drilling_intelligence.database.models import (
    Base,
    CasingRun,
    CementJob,
    CostItem,
    HseIncident,
    NptRecord,
    WellControlEvent,
    WellEvent,
)
from drilling_intelligence.operations.promote import VersionPromoter


def _derived_tables() -> list[tuple[str, object]]:
    """Every table the repository treats as source-versioned and derived - read from the ORM."""
    found = []
    for name, table in Base.metadata.tables.items():
        columns = {column.name for column in table.columns}
        if {"document_version_id", "origin"} <= columns:
            found.append((name, table))
    return sorted(found)


def _snapshot(workspace, table) -> list[tuple]:
    with workspace.database.read_only() as session:
        rows = session.execute(
            select(table).where(table.c.origin == KnowledgeOrigin.DERIVED.value)
        ).all()
    return sorted(tuple(str(value) for value in row) for row in rows)


def test_the_derived_table_inventory_is_taken_from_the_models(workspace) -> None:
    names = [name for name, _ in _derived_tables()]
    assert "well_control_event" in names
    assert "hse_incident" in names
    assert "cost_item" in names
    assert "casing_run" in names
    assert "cement_job" in names


def test_delete_promoted_removes_every_derived_row_the_version_owned(workspace) -> None:
    """The contract: un-derive this document version, and nothing it created survives."""
    ingest_v72(workspace)
    promote_file(workspace, "well_control_log_well-a3.xlsx")
    promote_file(workspace, "hse_register_well-a3.xlsx")

    with workspace.database.read_only() as session:
        version_id = str(session.execute(select(HseIncident.document_version_id)).scalars().first())

    with workspace.database.unit_of_work() as session:
        VersionPromoter(session).delete_promoted(version_id=version_id)

    leftovers = []
    with workspace.database.read_only() as session:
        for name, table in _derived_tables():
            count = len(
                session.execute(
                    select(table).where(
                        table.c.document_version_id == version_id,
                        table.c.origin == KnowledgeOrigin.DERIVED.value,
                    )
                ).all()
            )
            if count:
                leftovers.append((name, count))
    assert leftovers == [], f"derived rows survived delete_promoted: {leftovers}"


def test_delete_promoted_leaves_unrelated_rows_alone(workspace) -> None:
    """Exact snapshots, not counts: another document's rows must be byte-for-byte unchanged."""
    ingest_v72(workspace)
    promote_file(workspace, "well_control_log_well-a3.xlsx")
    promote_file(workspace, "hse_register_well-a3.xlsx")

    with workspace.database.read_only() as session:
        version_id = str(session.execute(select(HseIncident.document_version_id)).scalars().first())
    before = {name: _snapshot(workspace, table) for name, table in _derived_tables()}

    with workspace.database.unit_of_work() as session:
        VersionPromoter(session).delete_promoted(version_id=version_id)

    after = {name: _snapshot(workspace, table) for name, table in _derived_tables()}
    for name, _table in _derived_tables():
        if name == "hse_incident":
            continue
        assert before[name] == after[name], f"{name} was touched by an unrelated delete"
    assert before["hse_incident"] and not after["hse_incident"]


def test_delete_promoted_covers_the_older_domains_too(workspace) -> None:
    """The regression that motivated the fix: cost, casing and cement were being left behind."""
    ingest_v71(workspace)
    for file_name in ("casing_program_well-a3.xlsx", "cement_report_well-a3.xlsx"):
        try:
            promote_file(workspace, file_name)
        except Exception:  # noqa: BLE001 - a file the V7.1 corpus does not carry is simply skipped
            pass
    with workspace.database.read_only() as session:
        # Whatever the V7.1 corpus actually derived - casing, cement or cost - is the subject here.
        versions = []
        for model in (CostItem, CasingRun, CementJob):
            versions.extend(
                str(value)
                for value in session.execute(
                    select(model.document_version_id).where(
                        model.origin == KnowledgeOrigin.DERIVED.value
                    )
                ).scalars()
            )
    assert versions, "the V7.1 corpus produced no derived rows to test delete_promoted against"
    version_id = versions[0]

    with workspace.database.unit_of_work() as session:
        VersionPromoter(session).delete_promoted(version_id=version_id)

    with workspace.database.read_only() as session:
        for model in (CostItem, CasingRun, CementJob):
            table = model.__table__
            count = len(
                session.execute(
                    select(table).where(
                        table.c.document_version_id == version_id,
                        table.c.origin == KnowledgeOrigin.DERIVED.value,
                    )
                ).all()
            )
            assert count == 0, f"{model.__tablename__} survived delete_promoted"


def test_a_confirmed_row_survives_a_source_row_disappearing(workspace) -> None:
    """A human confirmation is not erased because an extraction changed underneath it."""
    ingest_v72(workspace)
    promote_file(workspace, "well_control_log_well-a3.xlsx")
    target = next(r for r in fetch(workspace, WellControlEvent) if r.event_label == "WC-01")

    with workspace.database.unit_of_work() as session:
        session.get(WellControlEvent, target.id).status = "CONFIRMED"

    from openpyxl import load_workbook

    path = workspace.root / "corpus" / "well_control_log_well-a3.xlsx"
    workbook = load_workbook(path)
    row = next(cell.row for cell in workbook.active["A"] if cell.value == "WC-01")
    workbook.active.delete_rows(row)
    workbook.save(path)
    reingest(workspace)
    promote_file(workspace, "well_control_log_well-a3.xlsx")

    # ``fetch`` returns history too, so the current answer has to be filtered for explicitly.
    current = {r.event_label: r for r in fetch(workspace, WellControlEvent) if r.is_current}
    assert "WC-01" in current, "a confirmed row must not silently disappear"
    assert current["WC-01"].status == "CONFIRMED"
    assert "WC-02" in current, "the untouched rows are still current"


def test_an_unconfirmed_row_the_source_stopped_stating_is_removed(workspace) -> None:
    ingest_v72(workspace)
    promote_file(workspace, "well_control_log_well-a3.xlsx")
    assert "WC-05" in {r.event_label for r in fetch(workspace, WellControlEvent)}

    from openpyxl import load_workbook

    path = workspace.root / "corpus" / "well_control_log_well-a3.xlsx"
    workbook = load_workbook(path)
    row = next(cell.row for cell in workbook.active["A"] if cell.value == "WC-05")
    workbook.active.delete_rows(row)
    workbook.save(path)
    reingest(workspace)
    promote_file(workspace, "well_control_log_well-a3.xlsx")

    assert "WC-05" not in {
        r.event_label for r in fetch(workspace, WellControlEvent) if r.is_current
    }


def test_no_cross_domain_row_is_created_or_lost_by_the_v72_lifecycle(workspace) -> None:
    ingest_v72(workspace)
    before = {
        name: _snapshot(workspace, table)
        for name, table in _derived_tables()
        if name not in ("well_control_event", "hse_incident")
    }
    promote_file(workspace, "well_control_log_well-a3.xlsx")
    promote_file(workspace, "hse_register_well-a3.xlsx")
    after = {
        name: _snapshot(workspace, table)
        for name, table in _derived_tables()
        if name not in ("well_control_event", "hse_incident")
    }
    assert before == after, "promoting V7.2 domains must not touch any other derived table"
    assert fetch(workspace, WellEvent) == []
    assert fetch(workspace, NptRecord) == []
