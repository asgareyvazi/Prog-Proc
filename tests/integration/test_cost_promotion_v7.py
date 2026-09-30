"""Cost promotion: source currency survives, planned and actual never merge, nothing is inferred.

Cost is the domain where a silent default is most expensive.  ``CostItem`` defaults both unit
columns to ``USD`` and ``currency_of("")`` folds an empty unit to ``USD``, so a writer that forgets
to read a currency does not fail - it restates a Norwegian kroner amount as dollars, at a factor of
about ten, in a field an engineer will later quote.  Every fixture here is therefore in a currency
other than USD wherever a currency matters.

The second thing these tests hold is that a cost line is not a licence to write a row.  A sheet that
mentions money, and a table whose money column does not say whether the figure was planned or spent,
are both refused; only a code-plus-description-plus-side-labelled-amount table becomes rows.
"""

from __future__ import annotations

from tests.fixtures.fieldops import fetch, ingest_v7, promote_file, reingest, well_id_for

from drilling_intelligence.database.models import CostItem, NptRecord

LEDGER = "cost_ledger_well-a3.xlsx"
MULTI = "cost_multi_currency_well-a3.xlsx"
NARRATIVE = "cost_narrative_well-a3.xlsx"


def _lines(workspace) -> dict[str, CostItem]:
    """Every cost line in the workspace, keyed by its source cost code."""
    return {row.cbs_code or row.wbs_code or "": row for row in fetch(workspace, CostItem)}


def _reasons(result) -> list[str]:
    return [str(item.get("reason") or "") for item in result.skipped]


def test_a_cost_ledger_becomes_one_row_per_line(workspace) -> None:
    ingest_v7(workspace)
    result = promote_file(workspace, LEDGER)
    assert result.outcome == "PROMOTED", result.to_dict()
    assert result.counts["cost_item"]["created"] == 4, result.to_dict()

    lines = _lines(workspace)
    assert set(lines) == {"1.2.4", "1.2.5", "1.2.6", "1.2.7"}

    rig = lines["1.2.4"]
    assert rig.description == "Rig - day rate, surface section"
    assert rig.planned_value == 1250000.0
    assert rig.actual_value == 1310500.0
    # The whole point of the fixture being in NOK: the column default is USD, so this assertion
    # fails the moment a writer stops reading the currency the source printed.
    assert rig.planned_unit == "NOK", rig.planned_unit
    assert rig.actual_unit == "NOK", rig.actual_unit
    assert rig.status == "CANDIDATE"
    assert rig.origin == "DERIVED"
    assert rig.record_state == "CURRENT"
    assert rig.provenance, "a promoted cost line must carry its source locator"
    assert rig.well_id, "the corpus file is filed under well A-3"


def test_planned_and_actual_stay_in_their_own_columns(workspace) -> None:
    """A line with a budget and no actual keeps NULL rather than borrowing the budget."""
    ingest_v7(workspace)
    promote_file(workspace, LEDGER)
    cement = _lines(workspace)["1.2.6"]
    assert cement.planned_value == 620000.0
    assert cement.actual_value is None, "an unstated actual must not be filled from the plan"


def test_a_forecast_column_is_neither_planned_nor_actual(workspace) -> None:
    """``Forecast`` is a prediction; folding it into planned or actual would assert a figure."""
    ingest_v7(workspace)
    result = promote_file(workspace, LEDGER)
    assert "UNMAPPED_MONEY_COLUMN" in _reasons(result), _reasons(result)

    rig = _lines(workspace)["1.2.4"]
    assert rig.planned_value == 1250000.0
    assert rig.actual_value == 1310500.0
    # 1 300 000 is the forecast.  It must not have displaced either figure.
    assert 1300000.0 not in {rig.planned_value, rig.actual_value}


def test_an_unrecognised_category_keeps_the_source_wording(workspace) -> None:
    """``Casing`` is not in the platform vocabulary; erasing it would lose the sheet's own meaning."""
    ingest_v7(workspace)
    promote_file(workspace, LEDGER)
    casing = _lines(workspace)["1.2.5"]
    assert casing.category != ""
    assert casing.attributes.get("source_wording", {}).get("category") == "Casing", (
        casing.attributes
    )


def test_two_currencies_in_one_table_are_not_conflated(workspace) -> None:
    """One table, planned in USD, actual in EUR: each side keeps its own currency."""
    ingest_v7(workspace)
    result = promote_file(workspace, MULTI)
    assert result.outcome == "PROMOTED", result.to_dict()

    lines = _lines(workspace)
    directional = lines["2.1.1"]
    assert directional.planned_unit == "USD", directional.planned_unit
    assert directional.actual_unit == "EUR", directional.actual_unit
    assert directional.planned_value == 180000.0
    assert directional.actual_value == 165000.0
    # No conversion is performed anywhere: the two figures stay the numbers the source printed.
    assert directional.actual_value != round(directional.planned_value * 0.92, 2)


def test_a_cost_document_with_no_cost_table_is_refused(workspace) -> None:
    """Prose quoting an amount, and a column headed ``Amount``, are not a cost table."""
    ingest_v7(workspace)
    result = promote_file(workspace, NARRATIVE)
    assert result.outcome == "UNSUPPORTED", result.to_dict()
    assert "NO_RECOGNISED_TABLE" in _reasons(result), _reasons(result)
    assert fetch(workspace, CostItem) == []


def test_repromoting_the_same_ledger_creates_nothing_new(workspace) -> None:
    """Identity is the content, so a second pass matches rather than accumulates."""
    ingest_v7(workspace)
    promote_file(workspace, LEDGER)
    before = {row.id for row in fetch(workspace, CostItem)}

    reingest(workspace)
    result = promote_file(workspace, LEDGER)
    after = {row.id for row in fetch(workspace, CostItem)}

    assert after == before, "a second promotion must not add rows"
    assert result.counts["cost_item"]["created"] == 0, result.to_dict()
    assert result.counts["cost_item"]["unchanged"] == 4, result.to_dict()


def test_npt_is_not_inferred_from_a_cost_line(workspace) -> None:
    """A cost line near an NPT record is not an attribution; ``npt_id`` stays unset."""
    ingest_v7(workspace)
    promote_file(workspace, LEDGER)

    # An NPT record exists on the same well, which is exactly the proximity that must not be read
    # as causation.
    lines = _lines(workspace)
    assert any(row.well_id for row in lines.values())
    for row in lines.values():
        assert row.npt_id is None, "npt_id must only come from an exact explicit attribution"


def test_a_missing_currency_is_refused_rather_than_defaulted_to_usd(workspace) -> None:
    """The critical adversarial case, exercised at the parser where the refusal happens."""
    from drilling_intelligence.operations.cost import cost_line_entries

    payload = {
        "tables": [
            {
                "table_id": "t",
                "sheet": "C",
                "provenance": {"locator": "sheet:C!A1"},
                "rows": [
                    ["CBS", "Description", "Budget", "Actual"],
                    ["9.9", "Unstated currency", "1000", "900"],
                ],
            }
        ]
    }
    entries = cost_line_entries(payload)
    assert len(entries) == 1
    entry = entries[0]
    assert entry.planned_text == "1000"
    assert entry.planned_value is None, "no value may be stored under an unstated currency"
    assert entry.planned_currency == ""


def test_a_code_less_cost_table_is_not_read(workspace) -> None:
    """Without a code two similar lines collapse to one identity, so the table is not admitted."""
    from drilling_intelligence.operations.cost import cost_line_entries

    payload = {
        "tables": [
            {
                "table_id": "t",
                "sheet": "C",
                "provenance": {"locator": "sheet:C!A1"},
                "rows": [
                    ["Description", "Budget (NOK)", "Actual (NOK)"],
                    ["First line", "100", "90"],
                    ["Second line", "200", "180"],
                ],
            }
        ]
    }
    assert cost_line_entries(payload) == []


def test_ambiguous_amount_grouping_is_not_guessed(workspace) -> None:
    """``1 250 000`` and ``12,5`` are refused, because either reading could be tenfold wrong."""
    from drilling_intelligence.operations.cost import cost_line_entries

    payload = {
        "tables": [
            {
                "table_id": "t",
                "sheet": "C",
                "provenance": {"locator": "sheet:C!A1"},
                "rows": [
                    ["CBS", "Description", "Budget (NOK)", "Actual (NOK)"],
                    ["9.1", "Space grouped", "1 250 000", "12,5"],
                    ["9.2", "Clean", "1250000", "1310500"],
                ],
            }
        ]
    }
    entries = {entry.cbs_code: entry for entry in cost_line_entries(payload)}
    assert entries["9.1"].planned_value is None
    assert entries["9.1"].actual_value is None
    assert entries["9.2"].planned_value == 1250000.0


def test_an_unresolved_npt_reference_leaves_the_link_unset(workspace) -> None:
    """``NptRecord`` has no human-readable code, so a sheet's reference cannot resolve to a row."""
    from drilling_intelligence.core.ids import new_id
    from drilling_intelligence.operations.promote import VersionPromoter

    ingest_v7(workspace)
    well_id = well_id_for(workspace, "A-3")
    with workspace.database.session() as session:
        promoter = VersionPromoter(session)

        npt_id, detail = promoter._cost_npt(reference="NPT-014", well_id=well_id)
        assert npt_id == ""
        assert "left unset rather than guessed" in detail

        # An exact stored identity is the one attribution that is not a guess.
        record = NptRecord(
            id=new_id("npt"),
            well_id=well_id,
            category="equipment_failure",
            description="fixture NPT event",
        )
        session.add(record)
        session.flush()

        exact, no_detail = promoter._cost_npt(reference=record.id, well_id=well_id)
        assert exact == record.id
        assert no_detail == ""

        # A reference naming an NPT record of a *different* well is not this well's attribution.
        other = NptRecord(
            id=new_id("npt"),
            well_id=well_id_for(workspace, "B-11"),
            category="equipment_failure",
            description="other well",
        )
        session.add(other)
        session.flush()
        cross, cross_detail = promoter._cost_npt(reference=other.id, well_id=well_id)
        assert cross == ""
        assert cross_detail


def _edit_ledger(workspace, cell: str, value: object) -> None:
    """Change one figure in the corpus ledger, as an operator correcting a sheet would."""
    from openpyxl import load_workbook

    path = workspace.root / "corpus" / LEDGER
    workbook = load_workbook(path)
    workbook.active[cell] = value
    workbook.save(path)


def _current(workspace) -> dict[str, CostItem]:
    """The current statement of each cost line, which is what a total must be built from."""
    rows = [row for row in fetch(workspace, CostItem) if row.is_current]
    return {row.cbs_code or "": row for row in rows}


def test_a_corrected_figure_does_not_double_count(workspace) -> None:
    """The defect this regression exists for: two CURRENT rows for one line, totalling both.

    Before cost rows were source-owned, a changed amount was a different identity and therefore a
    second row, and both stayed current.  A ledger totalling 2 060 500 in actuals came to 4 060 499
    after one figure was corrected - four lines, five current rows, one line counted at its old
    value and its new one.
    """
    ingest_v7(workspace)
    promote_file(workspace, LEDGER)
    before = _current(workspace)
    assert sum(row.actual_value or 0 for row in before.values()) == 2060500.0

    _edit_ledger(workspace, "E5", 1999999)  # the actual for CBS 1.2.4
    reingest(workspace)
    result = promote_file(workspace, LEDGER)
    assert result.outcome == "PROMOTED", result.to_dict()

    after = _current(workspace)
    assert set(after) == {"1.2.4", "1.2.5", "1.2.6", "1.2.7"}, "still four lines, not five"
    assert after["1.2.4"].actual_value == 1999999.0
    assert sum(row.actual_value or 0 for row in after.values()) == 2749999.0


def test_the_superseded_line_survives_as_history(workspace) -> None:
    """Standing a row down is not deleting it: the earlier statement stays readable."""
    ingest_v7(workspace)
    promote_file(workspace, LEDGER)
    _edit_ledger(workspace, "E5", 1999999)
    reingest(workspace)
    promote_file(workspace, LEDGER)

    every = fetch(workspace, CostItem)
    current = [row for row in every if row.is_current]
    superseded = [row for row in every if not row.is_current]

    # The whole previous version stands down, not only the line that changed: a version is one
    # statement of the ledger, and half-superseding it would leave two partial truths.
    assert len(current) == 4
    assert len(superseded) == 4
    assert {row.document_version_id for row in current}.isdisjoint(
        {row.document_version_id for row in superseded}
    )

    original = next(row for row in superseded if row.cbs_code == "1.2.4")
    assert original.actual_value == 1310500.0, "the original figure is still there"
    assert original.document_version_id, "history still names the version that produced it"
    assert original.provenance, "and still points at the source it came from"


def test_a_line_the_source_stopped_stating_is_removed_on_replace(workspace) -> None:
    """Replacement removes only what this source version no longer states."""
    ingest_v7(workspace)
    promote_file(workspace, LEDGER)
    assert set(_current(workspace)) == {"1.2.4", "1.2.5", "1.2.6", "1.2.7"}

    from openpyxl import load_workbook

    path = workspace.root / "corpus" / LEDGER
    workbook = load_workbook(path)
    workbook.active.delete_rows(8)  # CBS 1.2.7
    workbook.save(path)
    reingest(workspace)
    promote_file(workspace, LEDGER)

    assert set(_current(workspace)) == {"1.2.4", "1.2.5", "1.2.6"}


def test_a_manually_entered_line_is_never_touched_by_promotion(workspace) -> None:
    """A line a person typed belongs to no artefact, so no artefact may supersede or sweep it."""
    ingest_v7(workspace)
    promote_file(workspace, LEDGER)

    from drilling_intelligence.engineering.costs import CostRepository

    with workspace.database.unit_of_work() as session:
        row, created = CostRepository(session).record_item(
            description="Manually entered contingency",
            cbs_code="9.9.9",
            category="contingency",
            planned_value=50000,
            planned_unit="NOK",
        )
        assert created
        manual_id = row.id

    reingest(workspace)
    promote_file(workspace, LEDGER)

    every = {row.id: row for row in fetch(workspace, CostItem)}
    assert manual_id in every, "promotion must not delete a manual line"
    assert every[manual_id].is_current is True, "promotion must not stand a manual line down"
    assert every[manual_id].origin == "MANUAL"
    assert _current(workspace)["9.9.9"].planned_value == 50000.0


def test_an_exact_rerun_changes_nothing(workspace) -> None:
    """Promoting the same artefact twice is a no-op, not a second set of rows."""
    ingest_v7(workspace)
    promote_file(workspace, LEDGER)
    first = {row.id for row in fetch(workspace, CostItem)}

    result = promote_file(workspace, LEDGER)
    second = {row.id for row in fetch(workspace, CostItem)}

    assert second == first
    assert result.counts["cost_item"]["created"] == 0, result.to_dict()
    assert result.counts["cost_item"]["unchanged"] == 4, result.to_dict()
    assert all(row.is_current for row in fetch(workspace, CostItem))
