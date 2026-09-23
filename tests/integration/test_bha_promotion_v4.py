"""BHA promotion: a tally becomes an ordered assembly, and nothing more than the source stated.

Every test here drives the real pipeline - a generated workbook, a real extractor, a real promotion -
because a BHA contract that is only exercised against a hand-built payload dict proves nothing about
the thing it is deployed against.  The cases are chosen around the ways this domain actually fails:
a component that is really a footnote, an assembly attached to the wrong hole section, a source that
names another well, a re-tally that reorders the string, and a run that replaced an earlier run.
"""

from __future__ import annotations

from openpyxl import load_workbook
from sqlalchemy import select
from tests.fixtures.fieldops import fetch, ingest_v4, promote, promote_file, reingest

from drilling_intelligence.core.enums import KnowledgeRelationType
from drilling_intelligence.database.integrity import check_domain_identities
from drilling_intelligence.database.models import BhaComponent, BhaReport, KnowledgeRelation, Well
from drilling_intelligence.wells.repository import WellRepository

FILE = "bha_tally_well-a3.xlsx"


def _promoted(workspace) -> BhaReport:
    reports = fetch(workspace, BhaReport)
    assert len(reports) == 1, reports
    return reports[0]


def test_the_tally_becomes_one_run_with_its_components_in_source_order(workspace) -> None:
    ingest_v4(workspace)
    result = promote_file(workspace, FILE)
    assert result.outcome == "PROMOTED", result.to_dict()

    report = _promoted(workspace)
    assert report.bha_number == "14"
    assert report.component_count == 6
    assert report.top_depth_value == 9000.0 and report.top_depth_unit == "ft"
    assert report.bottom_depth_value == 10125.0 and report.bottom_depth_unit == "ft"
    assert report.assembly_description == "8 1/2 in BHA run 14"
    assert report.status == "CANDIDATE"
    assert report.origin == "DERIVED"
    assert report.is_current is True
    assert report.provenance, "a promoted parent must cite its source"

    # ``fetch`` orders by primary key, which is a content hash and carries no meaning; the string's
    # order is the ``sequence`` column, so that is what the assertion sorts on.
    components = sorted(fetch(workspace, BhaComponent), key=lambda row: row.sequence)
    assert [row.sequence for row in components] == [1, 2, 3, 4, 5, 6]
    assert [row.component_type for row in components] == [
        "DRILL_COLLAR",
        "STABILIZER",
        "MUD_MOTOR",
        "MWD_TOOL",
        "JAR",
        "HEAVY_WEIGHT_DRILL_PIPE",
    ]
    first = components[0]
    # The source's own words survive beside the canonical type, which is a convenience and not the
    # evidence.
    assert first.source_label == "Drill collar 6-1/4 x 2-13/16"
    assert first.manufacturer == "Grant Prideco"
    assert first.model == "DC-625"
    assert first.serial_number == "SN-1041"
    assert (first.od_value, first.od_unit, first.od_text) == (6.25, "in", "6.25")
    assert (first.id_value, first.id_unit) == (2.8125, "in")
    assert (first.length_value, first.length_unit) == (30.0, "ft")
    assert first.quantity == 8
    assert first.quality == "VALID"
    # Every component cites the row it came from.
    for row in components:
        assert row.provenance, f"component {row.sequence} cites no source"
        assert row.provenance[0]["source_row_index"] >= 0
        assert row.well_id == report.well_id
        assert row.bha_report_id == report.id
        assert row.document_version_id == report.document_version_id


def test_a_footnote_in_the_description_column_is_not_a_component(workspace) -> None:
    """The tally sheet ends with a note; a note is not hardware and must not enter the string."""
    ingest_v4(workspace)
    promote_file(workspace, FILE)
    labels = [row.source_label for row in fetch(workspace, BhaComponent)]
    assert "Assembly sketch filed separately" not in labels
    assert len(labels) == 6


def test_no_assembly_total_or_derived_dimension_is_computed(workspace) -> None:
    """The platform does not do the driller's arithmetic: what is absent stays absent."""
    ingest_v4(workspace)
    promote_file(workspace, FILE)
    report = _promoted(workspace)
    # The model has no total-length column at all, and no component's value is derived from another.
    assert not hasattr(report, "total_length_value")
    stabilizer = next(
        row for row in fetch(workspace, BhaComponent) if row.component_type == "STABILIZER"
    )
    # The fixture leaves the stabilizer's ID blank; it stays blank rather than being inferred from
    # the OD and a wall thickness nobody stated.  The column's own unit ("ID (in)") is retained
    # beside the empty value, because that is what the source printed for the column - what must
    # never appear is a *number* in a cell the source left empty.
    assert stabilizer.id_value is None
    assert stabilizer.id_text == ""
    hwdp = next(
        row for row in fetch(workspace, BhaComponent) if row.component_type == "HEAVY_WEIGHT_DRILL_PIPE"
    )
    # A component with no serial number keeps an empty one; nothing is generated.
    assert hwdp.serial_number == ""
    assert hwdp.quality == "VALID"


def test_a_component_type_is_only_ever_the_sources_own_words(workspace) -> None:
    ingest_v4(workspace)
    promote_file(workspace, FILE)
    for row in fetch(workspace, BhaComponent):
        if row.component_type:
            # The canonical name must be traceable to the description, not to a guess about it.
            assert row.source_label, row.component_type


def test_repromoting_the_same_artefact_is_a_no_op(workspace) -> None:
    ingest_v4(workspace)
    first = promote_file(workspace, FILE)
    assert first.counts["bha_component"]["created"] == 6
    assert first.counts["bha_report"]["created"] == 1

    second = promote_file(workspace, FILE)
    assert second.outcome == "UNCHANGED", second.to_dict()
    assert second.counts["bha_component"]["unchanged"] == 6
    assert second.counts["bha_component"]["created"] == 0
    assert second.counts["bha_report"]["unchanged"] == 1

    assert len(fetch(workspace, BhaReport)) == 1
    assert len(fetch(workspace, BhaComponent)) == 6


def test_a_component_the_source_drops_leaves_the_earlier_string_readable(workspace) -> None:
    """A re-tally that omits a component supersedes the earlier assembly; it does not delete it.

    The earlier run is engineering history: somebody ran that string, and the fact that a later tally
    no longer lists the stabilizer does not make the earlier record untrue.  So the older report and
    its six components stay in the database, stood down, and the current string is the new one.
    """
    ingest_v4(workspace)
    promote_file(workspace, FILE)
    source = workspace.root / "corpus" / FILE
    workbook = load_workbook(source)
    workbook["BHA Tally"].delete_rows(3)  # the stabilizer
    workbook.save(source)

    reingest(workspace)
    result = promote_file(workspace, FILE)
    assert result.outcome == "PROMOTED", result.to_dict()
    assert result.counts["bha_component"]["created"] == 5, result.to_dict()

    current = sorted(
        (row for row in fetch(workspace, BhaComponent) if row.is_current),
        key=lambda row: row.sequence,
    )
    assert "Stabilizer 6-1/2 in" not in [row.source_label for row in current]
    # The tally's own position is part of its meaning, so the components below the removed one move
    # up rather than leaving a gap in the string.
    assert [row.sequence for row in current] == [1, 2, 3, 4, 5]
    assert [row.component_type for row in current] == [
        "DRILL_COLLAR",
        "MUD_MOTOR",
        "MWD_TOOL",
        "JAR",
        "HEAVY_WEIGHT_DRILL_PIPE",
    ]

    reports = fetch(workspace, BhaReport)
    assert len(reports) == 2
    by_state = {row.is_current: row for row in reports}
    assert by_state[True].component_count == 5
    assert by_state[False].component_count == 6
    assert by_state[False].status == "SUPERSEDED"
    # The superseded components are still there, still attached to the superseded report.
    history = [row for row in fetch(workspace, BhaComponent) if not row.is_current]
    assert len(history) == 6
    assert "Stabilizer 6-1/2 in" in [row.source_label for row in history]
    assert all(row.status == "SUPERSEDED" for row in history)
    assert {row.bha_report_id for row in history} == {str(by_state[False].id)}


def test_a_re_tally_that_replaces_a_value_reports_the_change_rather_than_editing_it(workspace) -> None:
    """A confirmed component is a person's statement; a later extraction may not rewrite it."""
    from drilling_intelligence.operations.service import OperationalService

    ingest_v4(workspace)
    promote_file(workspace, FILE)
    component = next(
        row for row in fetch(workspace, BhaComponent) if row.component_type == "MUD_MOTOR"
    )
    with workspace.database.session() as session:
        outcome = OperationalService.for_workspace(workspace).set_status(
            "bha_component",
            str(component.id),
            "CONFIRMED",
            by="bha.engineer",
            reason="checked against the running sheet",
            session=session,
        )
        assert outcome["status"] == "CONFIRMED"
        session.commit()

    source = workspace.root / "corpus" / FILE
    workbook = load_workbook(source)
    workbook["BHA Tally"]["F4"] = 5.25  # the motor's OD, restated differently
    workbook.save(source)
    reingest(workspace)
    promote_file(workspace, FILE)

    kept = next(row for row in fetch(workspace, BhaComponent) if row.id == component.id)
    assert kept.status == "CONFIRMED"
    assert kept.od_value == 5.0, "a confirmed row is never overwritten by a re-extraction"
    assert kept.is_current is False


def test_a_source_that_names_another_well_is_refused_not_reattached(workspace) -> None:
    ingest_v4(workspace)
    source = workspace.root / "corpus" / FILE
    workbook = load_workbook(source)
    workbook["Summary"]["B3"] = "B-11"
    workbook.save(source)
    reingest(workspace)

    result = promote_file(workspace, FILE)
    assert result.outcome == "ERROR", result.to_dict()
    assert result.error == "WELL_SCOPE_CONFLICT"
    assert {item["reason"] for item in result.skipped} >= {"WELL_SCOPE_CONFLICT"}
    assert fetch(workspace, BhaReport) == []
    assert fetch(workspace, BhaComponent) == []


def test_a_source_that_names_another_field_is_refused(workspace) -> None:
    ingest_v4(workspace)
    source = workspace.root / "corpus" / FILE
    workbook = load_workbook(source)
    workbook["Summary"]["B4"] = "Brent"
    workbook.save(source)
    reingest(workspace)

    result = promote_file(workspace, FILE)
    assert result.outcome == "ERROR", result.to_dict()
    assert result.error == "WELL_SCOPE_CONFLICT"
    assert fetch(workspace, BhaReport) == []


def test_an_explicit_section_name_is_attached_and_an_ambiguous_one_is_not(workspace) -> None:
    def add_sections(*, sizes: tuple[float, ...], names: tuple[str, ...]) -> None:
        with workspace.database.session() as session:
            well = session.scalar(select(Well).where(Well.name == "A-3"))
            wells = WellRepository(session)
            for index, (size, name) in enumerate(zip(sizes, names, strict=True), start=1):
                wells.get_or_create_section(well, name, sequence=index, hole_size_in=size)
            session.commit()

    ingest_v4(workspace)
    add_sections(sizes=(12.25, 8.5), names=("12 1/4 in Surface", "8 1/2 in Intermediate"))
    result = promote_file(workspace, FILE)
    report = _promoted(workspace)
    assert report.section_resolution == "ATTRIBUTE"
    assert report.section_id, "an exact hole-size match is the one attribute the contract accepts"
    with workspace.database.read_only() as session:
        name = session.scalar(
            select(Well.name).where(Well.id == report.well_id)
        )
    assert name == "A-3"
    assert result.outcome in {"PROMOTED", "UNCHANGED"}, result.to_dict()

    # A second 8 1/2 in section makes the match ambiguous, and an ambiguous match stays NULL.
    add_sections(sizes=(8.5,), names=("8 1/2 in Lateral",))
    reingest(workspace)
    ambiguous = promote_file(workspace, FILE)
    assert {item["reason"] for item in ambiguous.skipped} >= {"AMBIGUOUS_SECTIONS"}
    assert ambiguous.outcome == "AMBIGUOUS", ambiguous.to_dict()


def test_a_section_the_source_names_but_the_well_does_not_have_stays_null(workspace) -> None:
    ingest_v4(workspace)
    result = promote_file(workspace, FILE)
    report = _promoted(workspace)
    assert report.section_id is None
    assert report.section_resolution == "UNMATCHED"
    assert {item["reason"] for item in result.skipped} >= {"SECTION_NOT_FOUND"}
    # Every component inherits the parent's unresolved section rather than inventing one.
    for component in fetch(workspace, BhaComponent):
        assert component.section_id is None


def test_a_new_run_is_a_new_row_and_the_earlier_run_stays_readable(workspace) -> None:
    ingest_v4(workspace)
    promote_file(workspace, FILE)
    source = workspace.root / "corpus" / FILE
    workbook = load_workbook(source)
    workbook["Summary"]["B5"] = "15"  # the next BHA run
    workbook.save(source)
    reingest(workspace)
    promote_file(workspace, FILE)

    reports = fetch(workspace, BhaReport)
    assert len(reports) == 2, reports
    by_number = {row.bha_number: row for row in reports}
    assert set(by_number) == {"14", "15"}
    # A new source version stands the older one down without deleting it.
    assert by_number["15"].is_current is True
    assert by_number["14"].is_current is False
    assert by_number["14"].status == "SUPERSEDED"
    assert len(fetch(workspace, BhaComponent, bha_report_id=str(by_number["14"].id))) == 6


def test_the_knowledge_graph_carries_the_assembly_and_its_components(workspace) -> None:
    ingest_v4(workspace)
    promote_file(workspace, FILE)
    report = _promoted(workspace)
    with workspace.database.read_only() as session:
        relations = list(
            session.scalars(
                select(KnowledgeRelation).where(
                    KnowledgeRelation.source_id.in_([report.id, str(report.well_id)])
                )
            )
        )
    kinds = {(row.relation, row.target_type) for row in relations}
    assert (KnowledgeRelationType.WELL_HAS_BHA.value, "bha_report") in kinds
    assert (KnowledgeRelationType.BHA_HAS_COMPONENT.value, "bha_component") in kinds
    component_edges = [
        row for row in relations if row.relation == KnowledgeRelationType.BHA_HAS_COMPONENT.value
    ]
    assert len(component_edges) == 6
    for edge in component_edges:
        assert edge.source_id == report.id


def test_a_promoted_assembly_leaves_the_domain_invariants_clean(workspace) -> None:
    ingest_v4(workspace)
    promote_file(workspace, FILE)
    with workspace.database.read_only() as session:
        problems = check_domain_identities(session)
    assert problems == [], [str(problem) for problem in problems]


def test_a_bha_narrative_produces_no_rows_and_says_why(workspace) -> None:
    """Same classification, different source shape: the writer is not entered."""
    from tests.fixtures.fieldops import register_wells, well_id_for
    from tests.fixtures.generate import build_v4_forensic_corpus

    from drilling_intelligence.ingestion.pipeline import IngestionPipeline

    register_wells(workspace)
    corpus = workspace.root / "corpus"
    build_v4_forensic_corpus(corpus)
    result = IngestionPipeline(
        settings=workspace.settings,
        workspace_root=workspace.root,
        database=workspace.database,
    ).run(root=corpus, well_id=well_id_for(workspace, "A-3"))
    assert result.ok, result.error
    outcome = promote_file(workspace, "bha_report_well-a3.txt")
    assert outcome.outcome == "UNSUPPORTED", outcome.to_dict()
    assert {item["reason"] for item in outcome.skipped} == {"NO_RECOGNISED_TABLE"}
    assert fetch(workspace, BhaReport) == []
    assert fetch(workspace, BhaComponent) == []


def test_a_parts_list_without_dimensions_is_not_an_assembly(workspace) -> None:
    """The contract needs a sizing column: an equipment list is not a bottom hole assembly."""
    from openpyxl import Workbook

    ingest_v4(workspace)
    source = workspace.root / "corpus" / FILE
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Summary"
    summary["A3"], summary["B3"] = "Well", "A-3"
    tally = workbook.create_sheet("BHA Tally")
    tally.append(["No", "Description", "Make"])
    tally.append([1, "Drill collar 6-1/4", "Grant Prideco"])
    workbook.save(source)
    reingest(workspace)

    outcome = promote_file(workspace, FILE)
    assert outcome.outcome == "UNSUPPORTED", outcome.to_dict()
    assert {item["reason"] for item in outcome.skipped} == {"NO_RECOGNISED_TABLE"}
    assert fetch(workspace, BhaReport) == []


def test_an_unlinked_document_promotes_nothing(workspace) -> None:
    ingest_v4(workspace)
    with workspace.database.session() as session:
        from tests.fixtures.fieldops import document_id_for

        from drilling_intelligence.database.models import Document

        document = session.get(Document, document_id_for(workspace, FILE))
        document.well_id = None
        session.commit()
    outcome = promote_file(workspace, FILE)
    assert outcome.outcome == "MISSING_WELL", outcome.to_dict()
    assert fetch(workspace, BhaReport) == []


def test_promoting_the_whole_workspace_twice_is_stable(workspace) -> None:
    ingest_v4(workspace)
    first = promote(workspace)
    second = promote(workspace)
    assert second["outcomes"]["promoted"] == 0, second["outcomes"]
    assert second["outcomes"]["unchanged"] >= first["outcomes"]["promoted"]
    assert second["outcomes"]["error"] == 0
    assert len(fetch(workspace, BhaComponent)) == first["counts"]["bha_component"]["created"]
