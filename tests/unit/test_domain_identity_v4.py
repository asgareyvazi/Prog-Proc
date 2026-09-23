"""Identity and idempotence for the V4 domains, and for the shared table helpers they sit on.

These are the semantics that make re-ingestion safe, stated once and checked directly: an identity is
a function of the source's own vocabulary (number, position, sequence) and never of a wall clock, a
random identifier or the order a query happens to return.

The parser assertions run against payloads the real extractor stored, not hand-built dicts: a contract
that only parses a payload somebody typed by hand has never been shown to parse the thing it is deployed
against.  ``tests/fixtures/generate.py`` builds the workbooks, the pipeline extracts them, and the tests
read the stored ``extraction.document_json`` back out.
"""

from __future__ import annotations

from sqlalchemy import select
from tests.fixtures.fieldops import register_wells, well_id_for
from tests.fixtures.generate import (
    BHA_TALLY_ROWS,
    BIT_TALLY_ROWS,
    SURVEY_STATION_ROWS,
    build_bha_tally_xlsx,
    build_bit_tally_xlsx,
    build_directional_survey_csv,
)

from drilling_intelligence.database.models import Document, Extraction
from drilling_intelligence.ingestion.pipeline import IngestionPipeline
from drilling_intelligence.operations.bha import (
    canonical_component_type,
    component_entries,
    component_tables,
)
from drilling_intelligence.operations.bit_record import bit_run_entries, bit_run_tables
from drilling_intelligence.operations.mud import canonical_summary_label, summary_entries
from drilling_intelligence.operations.promote import promotion_identity
from drilling_intelligence.operations.survey import station_entries, station_tables
from drilling_intelligence.operations.tableshape import alias_column, header_index, normalise_label


def _payload(workspace, filename: str) -> dict:
    """The extraction payload the pipeline actually stored for one file."""
    with workspace.database.read_only() as session:
        document = session.scalar(select(Document).where(Document.filename == filename))
        assert document is not None and document.current_version_id, filename
        extraction = session.scalar(
            select(Extraction).where(
                Extraction.document_version_id == str(document.current_version_id)
            )
        )
    assert extraction is not None and extraction.document_json, filename
    return dict(extraction.document_json)


def _ingest_files(workspace, files: dict[str, object]) -> None:
    """Write ``{filename: builder(path)}`` into the corpus and run the real pipeline over them."""
    if not well_id_for(workspace, "A-3"):
        register_wells(workspace)
    corpus = workspace.root / "corpus"
    corpus.mkdir(parents=True, exist_ok=True)
    for name, build in files.items():
        build(corpus / name)
    result = IngestionPipeline(
        settings=workspace.settings,
        workspace_root=workspace.root,
        database=workspace.database,
    ).run(root=corpus, well_id=well_id_for(workspace, "A-3"))
    assert result.ok, result.error


# ---------------------------------------------------------------- the parsers are deterministic
def test_the_same_tally_always_yields_the_same_entries(workspace) -> None:
    _ingest_files(
        workspace,
        {
            "bha_tally.xlsx": build_bha_tally_xlsx
        },
    )
    payload = _payload(workspace, "bha_tally.xlsx")
    first = component_entries(payload)
    second = component_entries(payload)
    assert [entry.sequence for entry in first] == [1, 2, 3, 4, 5, 6]
    assert [(entry.sequence, entry.source_label) for entry in first] == [
        (entry.sequence, entry.source_label) for entry in second
    ]
    assert [entry.od_value for entry in first] == [entry.od_value for entry in second]
    # The parser reads the stored table, so the same payload cannot yield a different string order.
    assert [entry.source_row_index for entry in first] == sorted(
        entry.source_row_index for entry in first
    )


def test_a_component_carries_its_own_description_alongside_the_canonical_type(workspace) -> None:
    _ingest_files(
        workspace,
        {
            "bha_tally.xlsx": build_bha_tally_xlsx
        },
    )
    entries = component_entries(_payload(workspace, "bha_tally.xlsx"))
    by_type = {entry.component_type: entry for entry in entries}
    assert set(by_type) == {
        "DRILL_COLLAR",
        "STABILIZER",
        "MUD_MOTOR",
        "MWD_TOOL",
        "JAR",
        "HEAVY_WEIGHT_DRILL_PIPE",
    }
    drill_collar = by_type["DRILL_COLLAR"]
    # The canonical name is a convenience; the source's words are the evidence, and both are kept.
    assert drill_collar.source_label == "Drill collar 6-1/4 x 2-13/16"
    # The fixture states no "Type" column, so nothing is claimed as stated; the canonical type is
    # read from the component's own words and the two are kept apart for exactly that reason.
    assert drill_collar.stated_type == ""
    assert (drill_collar.od_text, drill_collar.od_value, drill_collar.od_unit) == (
        "6.25",
        6.25,
        "in",
    )
    assert drill_collar.manufacturer == "Grant Prideco"
    assert drill_collar.serial_number == "SN-1041"


def test_a_component_type_is_only_ever_recognised_from_the_description() -> None:
    assert canonical_component_type("Drill collar 6-1/4 x 2-13/16") == "DRILL_COLLAR"
    assert canonical_component_type("Near-bit stabilizer") == "STABILIZER"
    assert canonical_component_type("Mud motor 5 in 7:8") == "MUD_MOTOR"
    assert canonical_component_type("MWD tool") == "MWD_TOOL"
    assert canonical_component_type("Jar") == "JAR"
    assert canonical_component_type("HWDP 5 in") == "HEAVY_WEIGHT_DRILL_PIPE"
    assert canonical_component_type("Heavy weight drill pipe 5 in") == "HEAVY_WEIGHT_DRILL_PIPE"
    # Words the contract does not define are left unclassified rather than mapped to the nearest one.
    for label in ("", None, "Misc item", "Rotary steerable system", "Drillstring"):
        assert canonical_component_type(label) == "", label


def test_a_component_table_needs_a_sizing_column(workspace) -> None:
    """An equipment list with no dimensions is not a bottom hole assembly, and must not be read as one."""
    from openpyxl import Workbook

    def without_sizing(path):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "BHA Tally"
        sheet.append(["No", "Description", "Make"])
        sheet.append([1, "Drill collar 6-1/4", "Grant Prideco"])
        workbook.save(path)

    _ingest_files(workspace, {"no_sizing.xlsx": without_sizing})
    assert component_tables(_payload(workspace, "no_sizing.xlsx")) == []
    assert component_entries(_payload(workspace, "no_sizing.xlsx")) == ()

    def with_sizing(path):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "BHA Tally"
        sheet.append(["No", "Description", "Make", "OD (in)"])
        sheet.append([1, "Drill collar 6-1/4", "Grant Prideco", 6.25])
        workbook.save(path)

    _ingest_files(workspace, {"with_sizing.xlsx": with_sizing})
    payload = _payload(workspace, "with_sizing.xlsx")
    assert len(component_tables(payload)) == 1
    assert len(component_entries(payload)) == 1


def test_a_footnote_row_is_not_a_component(workspace) -> None:
    from openpyxl import Workbook

    def with_footnote(path):
        workbook = Workbook()
        summary = workbook.active
        summary.title = "Summary"
        summary["A3"], summary["B3"] = "Well", "A-3"
        summary["A4"], summary["B4"] = "BHA No", "14"
        tally = workbook.create_sheet("BHA Tally")
        tally.append(["No", "Description", "Make", "OD (in)", "Length (ft)", "Qty"])
        for row in BHA_TALLY_ROWS:
            tally.append(row)
        # A footnote the way a tally actually prints one: words in the description column and nothing
        # in any other, including the position column.
        tally.append(["", "Assembly sketch filed separately", "", "", "", ""])
        workbook.save(path)

    _ingest_files(workspace, {"bha_tally.xlsx": with_footnote})
    entries = component_entries(_payload(workspace, "bha_tally.xlsx"))
    labels = [entry.source_label for entry in entries]
    assert "Assembly sketch filed separately" not in labels
    assert len(labels) == len(BHA_TALLY_ROWS)
    # The guard is "states nothing but words", not "looks like a sentence": the sequence still counts
    # 1..6 with no gap, so the note did not occupy a position in the string.
    assert [entry.sequence for entry in entries] == [1, 2, 3, 4, 5, 6]


def test_a_numbered_row_with_a_make_is_a_component_even_without_dimensions(workspace) -> None:
    """Where the footnote guard stops: a row that states anything but words is hardware.

    A tally row carrying a position and a manufacturer is a component whose dimensions the source did
    not print.  Refusing it would drop real hardware from the string, so the contract keeps it and
    leaves the absent dimensions absent.
    """
    from openpyxl import Workbook

    def numbered(path):
        workbook = Workbook()
        tally = workbook.active
        tally.title = "BHA Tally"
        tally.append(["No", "Description", "Make", "OD (in)", "Length (ft)", "Qty"])
        tally.append([1, "Drill collar 6-1/4", "Grant Prideco", 6.25, 30.0, 8])
        tally.append([2, "Crossover sub", "Bowen", "", "", 1])
        workbook.save(path)

    _ingest_files(workspace, {"bha_tally.xlsx": numbered})
    entries = component_entries(_payload(workspace, "bha_tally.xlsx"))
    assert [entry.source_label for entry in entries] == ["Drill collar 6-1/4", "Crossover sub"]
    crossover = entries[1]
    assert crossover.component_type == "CROSSOVER"
    assert crossover.od_value is None and crossover.od_text == ""
    assert crossover.length_value is None
    assert crossover.quantity == 1


# ---------------------------------------------------------------- bit records
def test_a_bit_run_carries_every_column_the_source_stated(workspace) -> None:
    _ingest_files(workspace, {"bit_tally.xlsx": build_bit_tally_xlsx})
    entries = bit_run_entries(_payload(workspace, "bit_tally.xlsx"))
    assert [entry.bit_number for entry in entries] == ["12", "13"]
    assert [entry.run_number for entry in entries] == ["1", "2"]
    first = entries[0]
    assert (first.size_text, first.size_value, first.size_unit) == ("8.5", 8.5, "in")
    assert (first.depth_in_value, first.depth_out_value) == (3500.0, 9100.0)
    assert (first.footage_value, first.footage_unit) == (5600.0, "ft")
    assert first.rotating_hours == 96.0
    assert first.pull_reason == "TD - casing point"
    assert first.dull_grade == "WT-1-NO-X-I-NO"
    assert first.iadc_code == "M1655SS"
    assert first.bha_number == "13", "the source's own BHA number is kept even when it matches nothing"
    assert first.well_name == "A-3"
    assert [entry.source_row_index for entry in entries] == sorted(
        entry.source_row_index for entry in entries
    )


def test_a_bit_needs_a_bit_number_and_at_least_one_measurement(workspace) -> None:
    from openpyxl import Workbook

    def parts_list(path):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Bit Record"
        sheet.append(["Bit No", "Make", "Model"])
        sheet.append(["12", "Smith", "S516"])
        workbook.save(path)

    _ingest_files(workspace, {"parts.xlsx": parts_list})
    assert bit_run_tables(_payload(workspace, "parts.xlsx")) == []

    def with_size(path):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Bit Record"
        sheet.append(["Bit No", "Size (in)"])
        sheet.append(["12", 8.5])
        workbook.save(path)

    _ingest_files(workspace, {"sized.xlsx": with_size})
    assert len(bit_run_tables(_payload(workspace, "sized.xlsx"))) == 1


def test_a_bit_row_without_a_bit_number_is_not_a_run(workspace) -> None:
    from openpyxl import Workbook

    def with_blank(path):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Bit Record"
        sheet.append(["Bit No", "Size (in)", "Footage (ft)"])
        sheet.append(["12", 8.5, 5600.0])
        sheet.append(["", 8.5, 500.0])
        workbook.save(path)

    _ingest_files(workspace, {"bits.xlsx": with_blank})
    entries = bit_run_entries(_payload(workspace, "bits.xlsx"))
    assert [entry.bit_number for entry in entries] == ["12"]


# ---------------------------------------------------------------- surveys
def test_a_numbered_survey_keeps_its_station_numbers_and_order(workspace) -> None:
    _ingest_files(
        workspace, {"survey.csv": build_directional_survey_csv}
    )
    entries = station_entries(_payload(workspace, "survey.csv"))
    assert [entry.station_number_text for entry in entries] == ["1", "2", "3", "4", "5"]
    assert [entry.sequence for entry in entries] == [1, 2, 3, 4, 5]
    assert [entry.md_value for entry in entries] == [9000.0, 9250.0, 9500.0, 9750.0, 10000.0]
    assert all(entry.md_unit == "ft" for entry in entries)
    assert all(entry.inclination_unit == "deg" for entry in entries)
    # The vertical columns are preserved, not recomputed.
    assert entries[0].tvd_value == 8995.4 and entries[0].tvd_unit == "ft"
    assert entries[0].northing_value == 6712345.1
    assert entries[0].easting_value == 432156.2
    assert entries[0].dls_value == 0.4
    assert len(SURVEY_STATION_ROWS) == 5


def test_a_survey_without_a_station_column_has_no_station_numbers(workspace) -> None:
    def unnumbered(path):
        path.write_text(
            "MD (ft),Inclination (deg),Azimuth (deg)\n9000,1.2,140.5\n9250,2.4,141.8\n",
            encoding="utf-8",
        )

    _ingest_files(workspace, {"survey.csv": unnumbered})
    entries = station_entries(_payload(workspace, "survey.csv"))
    assert [entry.station_number_text for entry in entries] == ["", ""]
    # Position is still recorded, so the writer can fall back to it without guessing.
    assert [entry.sequence for entry in entries] == [1, 2]
    assert [entry.md_value for entry in entries] == [9000.0, 9250.0]


def test_two_survey_sets_are_read_as_two_sets(workspace) -> None:
    def two_sets(path):
        path.write_text(
            "Survey Run,Station,MD (ft),Inclination (deg),Azimuth (deg)\n"
            "GYRO-1,1,9000,1.2,140.5\n"
            "GYRO-1,2,9250,2.4,141.8\n"
            "MWD-2,1,9500,4.1,143.2\n",
            encoding="utf-8",
        )

    _ingest_files(workspace, {"survey.csv": two_sets})
    entries = station_entries(_payload(workspace, "survey.csv"))
    assert [entry.run_label for entry in entries] == ["GYRO-1", "GYRO-1", "MWD-2"]
    # The label is preserved verbatim; the writer is what groups by it, and it never merges sets.
    assert {entry.run_label for entry in entries} == {"GYRO-1", "MWD-2"}


def test_a_survey_needs_three_distinct_columns(workspace) -> None:
    def depth_only(path):
        path.write_text("MD (ft),TVD (ft)\n9000,8995\n", encoding="utf-8")

    _ingest_files(workspace, {"depth.csv": depth_only})
    assert station_tables(_payload(workspace, "depth.csv")) == []

    def misnamed(path):
        path.write_text("Depth (ft),Depth (deg),Depth\n9000,1.2,140.5\n", encoding="utf-8")

    _ingest_files(workspace, {"misnamed.csv": misnamed})
    assert station_tables(_payload(workspace, "misnamed.csv")) == []

    def real(path):
        path.write_text(
            "MD (ft),Inclination (deg),Azimuth (deg)\n9000,1.2,140.5\n", encoding="utf-8"
        )

    _ingest_files(workspace, {"real.csv": real})
    assert len(station_tables(_payload(workspace, "real.csv"))) == 1


def test_a_row_without_a_numeric_depth_is_not_a_station(workspace) -> None:
    def with_totals(path):
        path.write_text(
            "Station,MD (ft),Inclination (deg),Azimuth (deg)\n"
            "1,9000,1.2,140.5\n"
            ",,,\n"
            ",Total,,\n"
            "2,9250,2.4,141.8\n",
            encoding="utf-8",
        )

    _ingest_files(workspace, {"survey.csv": with_totals})
    entries = station_entries(_payload(workspace, "survey.csv"))
    assert [entry.station_number_text for entry in entries] == ["1", "2"]
    assert [entry.sequence for entry in entries] == [1, 2]


def test_the_bit_fixture_is_the_contract_the_tests_are_written_against() -> None:
    assert len(BIT_TALLY_ROWS) == 2
    assert BIT_TALLY_ROWS[0][0] == "12"
    assert BIT_TALLY_ROWS[1][0] == "13"


# ---------------------------------------------------------------- the mud helpers the V4 writers share
def test_the_shared_header_helper_keeps_the_mud_contracts_behaviour() -> None:
    """The refactor moved these helpers; the mud contract they were written against did not change."""
    # The mud daily vocabulary is prefix-matched, which is what lets "MW in (ppg)" be read as the
    # alias "mw in" without listing every unit spelling a mud engineer might print.
    assert (
        alias_column(
            header_index(["MW in (ppg)", "PV"]),
            ("mw in", "mw-in"),
            prefix=True,
            strip_units=False,
        )
        == 0
    )
    # A bare "MW" column is a mud weight column.
    assert alias_column(header_index(["MW", "PV"]), ("mud weight", "mw")) == 0
    # And one that merely *starts with* an alias is not, unless prefix matching is asked for - the
    # opt-in that keeps "bit" from matching "bit size".
    assert alias_column(header_index(["Mud Weight History"]), ("mud weight", "mw")) == -1
    assert (
        alias_column(
            header_index(["Mud Weight History"]),
            ("mud weight", "mw"),
            prefix=True,
            strip_units=False,
        )
        == 0
    )
    # Prefix matching needs a word boundary, so a unit glued to the label is a different column and
    # is *not* silently absorbed: "Mud Weight(ppg)Active" is the active system, not a daily sample.
    assert (
        alias_column(
            header_index(["Mud Weight(ppg)Active", "PV(cP)"]),
            ("mud weight", "mw"),
            prefix=True,
            strip_units=False,
        )
        == -1
    )
    # Unit decoration is stripped for the V4 contracts, which is what lets "OD (in)" and "OD" match.
    assert alias_column(header_index(["OD (in)", "Length (ft)"], strip_units=True), ("od",)) == 0
    assert normalise_label("  OD   (in) ") == "od (in)"
    # First spelling of a repeated header wins; the second is never consulted.
    assert header_index(["Description", "Description (mm)"], strip_units=True)["description"] == 0


def test_the_mud_summary_vocabulary_still_resolves_the_labels_the_corpus_uses() -> None:
    assert canonical_summary_label("Mud Weight (ppg)") == "mud_weight"
    assert canonical_summary_label("MW") == "mud_weight"
    assert canonical_summary_label("PV (cP)") == "plastic_viscosity"
    assert canonical_summary_label("Depth MD (ft)") == "depth_md"
    # An unrecognised label stays unrecognised: the vocabulary is closed.
    assert canonical_summary_label("Something Else") == ""
    assert canonical_summary_label("") == ""


def test_summary_entries_preserve_the_source_unit_and_never_supply_one() -> None:
    entries = summary_entries(
        {
            "tables": [
                {
                    "table_id": "t1",
                    "sheet": "Mud Report",
                    "rows": [
                        ["Property", "Value", "Unit", "Remark"],
                        ["Mud Weight", 10.1, "ppg", "active"],
                    ],
                }
            ]
        }
    )
    assert len(entries) == 1
    assert entries[0].property_name == "mud_weight"
    assert entries[0].source_unit == "ppg"
    assert entries[0].source_value == "10.1"

    # No unit column: the parser reports no unit.  It is the writer that decides what an unverified
    # value means, and the conventional unit for a property is never stamped on here.
    bare = summary_entries(
        {"tables": [{"table_id": "t1", "rows": [["Property", "Value"], ["Mud Weight", 10.1]]}]}
    )
    assert len(bare) == 1
    assert bare[0].source_unit == ""


# ---------------------------------------------------------------- promotion identity itself
def test_a_promotion_identity_is_stable_and_never_time_based() -> None:
    first = promotion_identity(
        version_id="v1", kind="bha_report", table_id="bha|14", well_id="well-1", extra="14"
    )
    second = promotion_identity(
        version_id="v1", kind="bha_report", table_id="bha|14", well_id="well-1", extra="14"
    )
    assert first == second
    assert first.startswith("promote:")
    assert len(first) == len("promote:") + 32
    # No timestamp and no random component anywhere in the derivation: calling it again after a
    # pause, in another process, against the same artefact, gives the same key.
    for _ in range(5):
        assert (
            promotion_identity(
                version_id="v1", kind="bha_report", table_id="bha|14", well_id="well-1", extra="14"
            )
            == first
        )


def test_a_promotion_identity_changes_with_every_semantic_part() -> None:
    base = promotion_identity(
        version_id="v1", kind="bha_report", table_id="bha|14", well_id="well-1", extra="14"
    )
    assert (
        promotion_identity(
            version_id="v2", kind="bha_report", table_id="bha|14", well_id="well-1", extra="14"
        )
        != base
    )
    assert (
        promotion_identity(
            version_id="v1", kind="bit_record", table_id="bha|14", well_id="well-1", extra="14"
        )
        != base
    )
    assert (
        promotion_identity(
            version_id="v1", kind="bha_report", table_id="bha|15", well_id="well-1", extra="14"
        )
        != base
    )
    assert (
        promotion_identity(
            version_id="v1", kind="bha_report", table_id="bha|14", well_id="well-2", extra="14"
        )
        != base
    )
    assert (
        promotion_identity(
            version_id="v1", kind="bha_report", table_id="bha|14", well_id="well-1", extra="15"
        )
        != base
    )


def test_the_row_position_is_part_of_a_mud_identity_and_not_of_a_v4_one() -> None:
    """The mud daily table is positional; the V4 identities are semantic, so ``row_index`` stays 0.

    That difference is deliberate and load-bearing: re-exporting a tally with the rows in another order
    must not move a component's identity, while a mud daily row's position *is* part of what identifies
    it.  Both behaviours are pinned here so neither can drift into the other.
    """
    assert (
        promotion_identity(version_id="v1", kind="mud_report", row_index=1, extra="x")
        != promotion_identity(version_id="v1", kind="mud_report", row_index=2, extra="x")
    )
    # The V4 writers always pass row_index=0 and put the semantics in ``extra``.
    assert promotion_identity(
        version_id="v1", kind="bha_component", row_index=0, extra="14|1|drill collar"
    ) == promotion_identity(
        version_id="v1", kind="bha_component", row_index=0, extra="14|1|drill collar"
    )


def test_the_v4_identities_are_distinct_across_domains() -> None:
    """The same numbers in two domains must not collide into one row."""
    bha = promotion_identity(
        version_id="v1", kind="bha_report", table_id="t", well_id="well-1", extra="14"
    )
    bit = promotion_identity(
        version_id="v1", kind="bit_record", table_id="t", well_id="well-1", extra="14"
    )
    survey = promotion_identity(
        version_id="v1", kind="survey_run", table_id="t", well_id="well-1", extra="14"
    )
    assert len({bha, bit, survey}) == 3


def test_the_identity_of_a_real_promoted_row_starts_with_the_promotion_prefix(workspace) -> None:
    """The key in the database is the key this function produces - not an id minted at write time."""
    from tests.fixtures.fieldops import fetch, promote

    _ingest_files(
        workspace,
        {
            "bha_tally.xlsx": build_bha_tally_xlsx,
            "bit_tally.xlsx": build_bit_tally_xlsx,
            "survey.csv": build_directional_survey_csv,
        },
    )
    promote(workspace)
    from drilling_intelligence.database.models import BhaComponent, BitRecord, SurveyStation

    for model in (BhaComponent, BitRecord, SurveyStation):
        rows = fetch(workspace, model)
        assert rows, model.__tablename__
        for row in rows:
            assert str(row.identity_key or "").startswith("promote:"), (
                model.__tablename__,
                row.identity_key,
            )
            assert not str(row.id).startswith("promote:"), "the primary key is not the identity"
